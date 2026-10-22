"""Build the Word manuscript (Manuscript_MolInf.docx) and the cover letter from paper/AChE_MolInf_v4/main.tex, so the Word and LaTeX versions cannot disagree.
Handles the LaTeX subset used in main.tex (sections, paragraphs, \\textit/\\textbf/\\textsuperscript, simple math, \\cite, tables, figures, bibliography).
Run with bbbenv python (python-docx) from the project root after compiling main.tex (needs main.aux and main.bbl)."""
import re, os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

D = "paper/AChE_MolInf_v4/"
tex = open(D + "main.tex", encoding="utf-8").read()
aux = open(D + "main.aux", encoding="utf-8").read()
bbl = open(D + "main.bbl", encoding="utf-8").read()
num = {k: int(n) for k, n in re.findall(r"\\bibcite\{([^}]+)\}\{\{(\d+)\}", aux)}
BS = "\\"
SYM = [(BS + "geq", "\u2265"), (BS + "leq", "\u2264"), (BS + "cdot", "\u00b7"), (BS + "alpha", "\u03b1"), (BS + "beta", "\u03b2"), (BS + "mu", "\u00b5"), (BS + "rho", "\u03c1"), (BS + "pi", "\u03c0"),
       ("$-$", "\u2212"), (BS + "AA{}", "\u00c5"), (BS + "AA", "\u00c5"), ("--", "\u2013"), ("~", "\u00a0"), (BS + "%", "%"), (BS + "&", "&"), (BS + "_", "_"), ("``", "\u201c"), ("''", "\u201d"),
       (BS + ",", " "), (BS + "textmu{}", "\u00b5"), ("\\textquotesingle", "'")]


def strip_math(s):
    s = re.sub(r"\$([^$]*)\$", lambda m: m.group(1), s)
    return s


def clean_text(s):
    for a, b in SYM:
        s = s.replace(a, b)
    s = s.replace("{\\ttfamily\\scriptsize\\seqsplit{", "").replace("\\seqsplit{", "")
    s = re.sub(r"\\hyperref\[[^\]]*\]\{([^}]*)\}", r"\1", s)
    s = re.sub(r"\\url\{([^}]*)\}", r"\1", s)
    return s


def cite_repl(m):
    keys = [k.strip() for k in m.group(1).split(",")]
    ns = sorted(num[k] for k in keys if k in num)
    return "[" + ", ".join(str(n) for n in ns) + "]"


# tokens of inline formatting
def add_runs(par, s, size=None, base_bold=False, base_italic=False):
    s = re.sub(r"\\cite\{([^}]*)\}", cite_repl, s)
    s = re.sub(r"\\ref\{[^}]*\}", "?", s) if False else s
    s = clean_text(s)
    s = strip_math(s)
    pos = 0
    pat = re.compile(r"\\textit\{((?:[^{}]|\{[^{}]*\})*)\}|\\textbf\{((?:[^{}]|\{[^{}]*\})*)\}|\\textsuperscript\{((?:[^{}]|\{[^{}]*\})*)\}|\\textsubscript\{((?:[^{}]|\{[^{}]*\})*)\}")
    for m in pat.finditer(s):
        if m.start() > pos:
            r = par.add_run(re.sub(r"[{}]", "", s[pos:m.start()])); r.bold = base_bold; r.italic = base_italic
            if size: r.font.size = Pt(size)
        txt, kind = next((g, i) for i, g in enumerate(m.groups()) if g is not None)
        r = par.add_run(re.sub(r"[{}]", "", txt))
        r.bold = base_bold or kind == 1; r.italic = base_italic or kind == 0
        if kind == 2: r.font.superscript = True
        if kind == 3: r.font.subscript = True
        if size: r.font.size = Pt(size)
        pos = m.end()
    if pos < len(s):
        r = par.add_run(re.sub(r"[{}]", "", s[pos:])); r.bold = base_bold; r.italic = base_italic
        if size: r.font.size = Pt(size)


doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.5); s.top_margin = s.bottom_margin = Cm(2.5)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color); tcPr.append(shd)


# ---- front matter
doc.add_paragraph().add_run("Full Paper").italic = True
p = doc.add_paragraph(); add_runs(p, re.search(r"\{\\large\\bfseries ([^\\]+?)\\par\}", tex).group(1), size=15, base_bold=True)
p = doc.add_paragraph("Uday Raj, Mahim Aggarwal, Ishank Sharma, and Chandra Mohan Dasari*")
p = doc.add_paragraph("Indian Institute of Information Technology Design and Manufacturing Kurnool, Jagannathagattu Hill, Kurnool, Andhra Pradesh 518008, India. E-mail: chandramohan.d@iiitk.ac.in")
for r in p.runs: r.font.size = Pt(9)
doc.add_heading("Abstract", 2)
ab = re.search(r"\\section\*\{Abstract\}\s*\{\\small (.*?)\\par\}", tex, re.S).group(1)
add_runs(doc.add_paragraph(), ab)
kw = re.search(r"\\textbf\{Keywords:\} (.*?)\\par", tex, re.S).group(1)
p = doc.add_paragraph(); r = p.add_run("Keywords: "); r.bold = True; add_runs(p, kw.replace(BS + "cdot", ";"))
doc.add_page_break()

# ---- body
body = tex[tex.index(BS + "section{Introduction}"):tex.index(BS + "section*{Statement of Contribution}")]
sec_n, sub_n = 0, 0
fig_n, tab_n = 0, 0
blocks = re.split(r"(\\begin\{table\}.*?\\end\{table\}|\\begin\{figure\}.*?\\end\{figure\})", body, flags=re.S)


def emit_text(chunk):
    global sec_n, sub_n
    for para in re.split(r"\n\s*\n", chunk):
        para = para.strip()
        para = "\n".join(ln for ln in para.splitlines() if not ln.lstrip().startswith("%")).strip()   # drop %-comment lines (marker lines of inserted blocks) but keep the text under them
        if not para:
            continue
        m = re.match(r"\\section\{([^}]*)\}", para)
        if m:
            sec_n += 1; sub_n = 0; doc.add_heading("%d. %s" % (sec_n, m.group(1)), 1); para = para[m.end():].strip()
            para = re.sub(r"^\\label\{[^}]*\}\s*", "", para)
            if not para: continue
        m = re.match(r"\\subsection\{([^}]*)\}", para)
        if m:
            sub_n += 1; doc.add_heading("%d.%d. %s" % (sec_n, sub_n, m.group(1)), 2); para = para[m.end():].strip()
            para = re.sub(r"^\\label\{[^}]*\}\s*", "", para)
            if not para: continue
        para = re.sub(r"\\label\{[^}]*\}", "", para)
        add_runs(doc.add_paragraph(), para)


labels = {}

# section / subsection numbers for \ref{sec:...}
_sn = _sb = 0
_pos = [(m_.start(), m_.group(1), m_.group(2)) for m_ in re.finditer(r"\\(section|subsection)\{[^}]*\}(?:\s*\\label\{([^}]*)\})?", body)]
for _, kind, lab in _pos:
    if kind == "section": _sn += 1; _sb = 0
    else: _sb += 1
    if lab and lab.startswith("sec:"):
        labels[lab] = str(_sn) if kind == "section" else "%d.%d" % (_sn, _sb)
# first pass: numbering of tables/figures in order of appearance
tn = fn = 0
for blk in blocks:
    if blk.startswith(BS + "begin{table}"):
        tn += 1; labels["tab:" + re.search(r"\\label\{tab:([^}]*)\}", blk).group(1)] = tn
    elif blk.startswith(BS + "begin{figure}"):
        fn += 1; labels["fig:" + re.search(r"\\label\{fig:([^}]*)\}", blk).group(1)] = fn
_orig_add_runs = add_runs


def add_runs2(par, s, **kw):
    s = re.sub(r"\\'([aeiouAEIOU])", lambda m: {'a': 'á', 'e': 'é', 'i': 'í', 'o': 'ó', 'u': 'ú', 'A': 'Á', 'E': 'É', 'I': 'Í', 'O': 'Ó', 'U': 'Ú'}[m.group(1)], s)
    s = re.sub(r"\\ref\{(sec:[^}]*)\}", lambda m: str(labels.get(m.group(1), '?')), s)
    s = re.sub(r"#(tab|fig)([A-Za-z0-9_:.]+)", lambda m: ("Table " if m.group(1) == "tab" else "Figure ") + str(labels.get(m.group(1) + ":" + m.group(2), "?")), s)
    s = re.sub(r"\\ref\{(tab|fig):([^}]*)\}", lambda m: str(labels.get(m.group(1) + ":" + m.group(2), "?")), s)
    return _orig_add_runs(par, s, **kw)


add_runs = add_runs2
for blk in blocks:
    if blk.startswith(BS + "begin{table}"):
        n = labels["tab:" + re.search(r"\\label\{tab:([^}]*)\}", blk).group(1)]
        cap = re.search(r"\\caption\{(.*?)\}\n\\label", blk, re.S).group(1)
        p = doc.add_paragraph(); r = p.add_run("Table %d. " % n); r.bold = True; add_runs(p, cap, size=9.5)
        rows = []
        tab = re.search(r"\\begin\{tabular\}.*?\\toprule(.*?)\\bottomrule", blk, re.S).group(1)
        tab = tab.replace(BS + "midrule", "")
        for line in re.split(r"\\\\\s*\n", tab):
            line = line.strip()
            if not line or line.startswith("%"): continue
            rows.append([c.strip() for c in line.split("&")])
        ncol = max(len(r) for r in rows)
        t = doc.add_table(rows=0, cols=ncol); t.style = "Table Grid"
        for i, row in enumerate(rows):
            cells = t.add_row().cells
            for j in range(ncol):
                cells[j].text = ""
                add_runs(cells[j].paragraphs[0], row[j] if j < len(row) else "", size=8.5, base_bold=(i == 0))
                if i == 0: shade(cells[j], "D9E2F3")
        fn_m = re.search(r"\\begin\{minipage\}\{\\textwidth\}(.*?)\\end\{minipage\}", blk, re.S)
        if fn_m: add_runs(doc.add_paragraph(), fn_m.group(1).replace(BS + "footnotesize", ""), size=8.5)
    elif blk.startswith(BS + "begin{figure}"):
        n = labels["fig:" + re.search(r"\\label\{fig:([^}]*)\}", blk).group(1)]
        img = re.search(r"\\includegraphics\[[^\]]*\]\{figures/([^}]*)\.pdf\}", blk).group(1)
        path = D + "figures/" + img + ".png"
        if os.path.exists(path):
            doc.add_picture(path, width=Cm(15.5)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = re.search(r"\\caption\{(.*?)\}\n\\label", blk, re.S).group(1)
        p = doc.add_paragraph(); r = p.add_run("Figure %d. " % n); r.bold = True; add_runs(p, cap, size=9.5)
    else:
        emit_text(blk)

# ---- back matter
tail = tex[tex.index(BS + "section*{Statement of Contribution}"):tex.index(BS + "bibliographystyle")]
for m in re.finditer(r"\\section\*\{([^}]*)\}\s*(.*?)(?=\\section\*|\Z)", tail, re.S):
    doc.add_heading(m.group(1), 2)
    for para in re.split(r"\n\s*\n", m.group(2).strip()):
        if para.strip(): add_runs(doc.add_paragraph(), para.strip())
doc.add_heading("References", 2)
for key, txt in re.findall(r"\\bibitem\{([^}]+)\}(.*?)(?=\\bibitem|\\end\{thebibliography\})", bbl, re.S):
    txt = re.sub(r"\\href\{[^}]*\}\{doi:~\\nolinkurl\{([^}]*)\}\}", r"doi: \1", txt)
    txt = re.sub(r"\\url\{([^}]*)\}", r"\1", txt)
    txt = re.sub(r"\s+", " ", txt).strip()
    p = doc.add_paragraph(); r = p.add_run("[%d] " % num[key]); add_runs(p, txt, size=9)
doc.save(D + "Manuscript_MolInf.docx")
print("Word manuscript:", D + "Manuscript_MolInf.docx", "| tables", tn, "figures", fn, "| references", len(num))

# ---- cover letter
cl = Document()
for s in cl.sections: s.left_margin = s.right_margin = Cm(2.5)
cl.styles["Normal"].font.name = "Times New Roman"; cl.styles["Normal"].font.size = Pt(11)
for line in ["[Date]", "", "The Editors, Molecular Informatics (Wiley-VCH)", "", "Dear Editors,", ""]:
    cl.add_paragraph(line)
paras = [
    "We submit the Full Paper \"Computational Design of 106 Novel, Brain-Penetrant Acetylcholinesterase Inhibitor Candidates with Computed Synthetic Routes\" for consideration in Molecular Informatics.",
    "Computational AChE papers usually stop at a predicted potency. Here 383,752 one- and two-edit analogues of 2,225 potent published inhibitors are reduced at the design stage to 2,700 designs, which are carried through activity, brain-penetration, drug-likeness, toxicity, docking, protonation, tautomer and stereochemistry gates, a four-source novelty search (PubChem exact and same-connectivity, ChEMBL flexmatch, SureChEMBL) and a computed retrosynthesis to purchasable starting materials. 106 molecules clear every gate; each has a computed route of one to six steps, and a synthesis dossier lists structure, route, predicted values with margins and the flags that decide what to measure first.",
    "What we think is of interest to readers of this journal: the models behind every gate are reported with held-out-scaffold performance and simple baselines, including the places where a component is weak (the docking score separates binders from matched decoys only weakly; the P-gp term did not separate BBB-permeable from non-permeable compounds; the activity model matches nearest-neighbor lookup on chemistry unlike its training data, a result we report together with the in-domain result, where it ranks later-published compounds clearly better than the baselines). We also anchor the design to measured data: the substitutions used are tested in real AChE compound pairs, each molecule is tied to a measured potent parent, and a blind re-discovery test, which does not favour our edit library over random edits, is reported as it came out. An independent audit script re-reads every gate from the delivered file, and it found, and we corrected, a gate that had not been applied to part of the candidate pool. A novelty search found nine otherwise eligible candidates in a series whose parents are covered by earlier patent and literature disclosures, and we excluded them.",
    "The paper also reports a priority tier of 63 molecules (margin of 10 % or more on every gate) and a shortlist of 20 for a first synthesis batch, chosen by ranking under 5,000 random weightings of eight checks so that the choice does not hinge on one set of weights. Because no assay data exist, the Discussion states falsifiable predictions for the first experiment: in particular, that molecules which keep the Trp86 interaction should be markedly more potent than those that lose it, a point on which the ligand-based models and the structure-based reading disagree.",
    "We state plainly that none of the molecules has been synthesized or assayed, that every property is a prediction, and that novelty and freedom to operate need expert assessment. The work has not been published elsewhere and is not under consideration by another journal. All authors have approved the submission.",
    "Suggested referees: [to be completed by the authors]. Excluded referees: [to be completed by the authors].",
    "Sincerely,", "Chandra Mohan Dasari (corresponding author), on behalf of all authors", "chandramohan.d@iiitk.ac.in"]
for t_ in paras: cl.add_paragraph(t_)
cl.save(D + "Cover_Letter.docx")
print("cover letter written")
